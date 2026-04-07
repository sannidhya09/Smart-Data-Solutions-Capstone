#!/usr/bin/env python3
"""Generate the Red Duke Technical Documentation Word document."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
import os

doc = Document()

# ── Page margins ──
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1.2)

# ── Style setup ──
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.15

# Heading styles
for level in [1, 2, 3]:
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'Calibri'
    hs.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    if level == 1:
        hs.font.size = Pt(20)
        hs.font.bold = True
    elif level == 2:
        hs.font.size = Pt(14)
        hs.font.bold = True
    elif level == 3:
        hs.font.size = Pt(12)
        hs.font.bold = True


def add_normal(text):
    """Add a normal paragraph."""
    p = doc.add_paragraph(text)
    return p


def add_blue(text):
    """Add a blue paragraph for non-developer sections."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(0x00, 0x50, 0xA0)
    return p


def add_bullet(text, blue=False):
    """Add a bullet point."""
    p = doc.add_paragraph(style='List Bullet')
    if blue:
        p.clear()
        run = p.add_run(text)
        run.font.color.rgb = RGBColor(0x00, 0x50, 0xA0)
    else:
        p.clear()
        p.add_run(text)
    return p


def add_code_block(text):
    """Add a monospaced code-like paragraph with gray background feel."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    p.paragraph_format.left_indent = Inches(0.4)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_separator():
    """Add a thin line separator."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run('_' * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)


# ═══════════════════════════════════════════════════════════════════════
# TITLE PAGE
# ═══════════════════════════════════════════════════════════════════════

doc.add_paragraph()  # spacer
doc.add_paragraph()

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Red Duke Intelligence Platform')
run.font.size = Pt(26)
run.font.bold = True
run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Technical Documentation')
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0x5B, 0x35, 0xC4)

doc.add_paragraph()

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = meta.add_run('Data Flow & Analysis Engine')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()
doc.add_paragraph()

info = doc.add_paragraph()
info.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = info.add_run('Author: Sannidhya Tiwari\nSmart Data Solutions\nMarch 2026')
run.font.size = Pt(11)
run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════
# TABLE OF CONTENTS (manual)
# ═══════════════════════════════════════════════════════════════════════

doc.add_heading('Table of Contents', level=1)
toc_items = [
    'Part 1: Data Flow — Developer Version',
    'Part 1: Data Flow — Non-Developer Version',
    'Part 2: Analysis Engine — Developer Version',
    'Part 2: Analysis Engine — Non-Developer Version',
]
for i, item in enumerate(toc_items, 1):
    p = doc.add_paragraph(f'{i}.  {item}')
    p.paragraph_format.space_after = Pt(3)

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════
# PART 1: DATA FLOW — DEVELOPER VERSION
# ═══════════════════════════════════════════════════════════════════════

doc.add_heading('Part 1: Data Flow', level=1)
doc.add_heading('Developer Version', level=2)

add_normal(
    'This section walks through how data moves through the Red Duke platform, '
    'from the moment documents sit in a client\'s SharePoint folder to the point '
    'where results appear in the React dashboard.'
)

# Step 0
doc.add_heading('Step 0 — SharePoint Download', level=3)
add_normal(
    'The pipeline starts in sharepoint_files.py. We authenticate against Microsoft Graph '
    'using a client-credentials OAuth2 flow (tenant ID, client ID, client secret stored in .env). '
    'The module calls the Graph API to resolve the SharePoint site ID and drive ID, then lists '
    'all files under the configured folder path (SP_FOLDER_PATH).'
)
add_normal(
    'Each file is downloaded via the Graph API\'s /content endpoint and saved to a local '
    'directory (defaults to red_duke_demo/downloads/). If SharePoint is not configured or '
    'the download fails, main.py falls back to the local sample_data/ folder — this makes '
    'the pipeline work in both production and demo scenarios without code changes.'
)

# Step 1
doc.add_heading('Step 1 — Document Discovery', level=3)
add_normal(
    'The discover_files() function in main.py walks the source directory recursively using '
    'os.walk() and collects every file whose extension matches our SUPPORTED_EXTENSIONS set:'
)
add_code_block(
    'SUPPORTED_EXTENSIONS = {\n'
    '    ".docx", ".xlsx", ".xlsm", ".pptx", ".pdf", ".vsdx",\n'
    '    ".csv", ".tsv", ".txt", ".md", ".png", ".jpg", ".jpeg"\n'
    '}'
)
add_normal(
    'Files are sorted alphabetically so the pipeline output is deterministic across runs.'
)

# Step 2
doc.add_heading('Step 2 — Document Parsing', level=3)
add_normal(
    'Each discovered file is routed to a type-specific parser based on its extension. '
    'Every parser returns a ParseResult dataclass containing:'
)
add_bullet('filename, filepath, file_type, file_size')
add_bullet('text — the full extracted text content')
add_bullet('sections — a list of {title, content} dicts (one per heading/sheet/slide)')
add_bullet('tables — extracted table data as lists of rows')
add_bullet('metadata — parser-specific info (author, sheet names, slide count, etc.)')
add_bullet('success flag and error message if parsing failed')

add_normal('The parsers work as follows:')

add_bullet('.docx — Uses python-docx. Iterates paragraphs, splits on heading styles to build sections. Extracts all tables as row arrays.')
add_bullet('.xlsx/.xlsm — Uses openpyxl. Reads every sheet, extracts headers and data rows, tracks cell counts.')
add_bullet('.pptx — Uses python-pptx. Reads each slide\'s text frames, tables, and speaker notes.')
add_bullet('.vsdx — Treats the file as a ZIP archive, parses the XML inside visio/pages/pageN.xml, and extracts Text elements from shapes. Builds a flow narrative by joining shape texts with arrows.')
add_bullet('.pdf — Uses PyMuPDF (fitz) to extract text from each page.')
add_bullet('.csv/.tsv — Read directly as text with appropriate delimiter.')
add_bullet('.txt/.md — Read as raw text.')

add_normal(
    'If any parser throws an exception, the ParseResult is still created but with success=False '
    'and the error message captured. The pipeline continues — one bad file does not stop the run.'
)

# Step 3
doc.add_heading('Step 3 — Chunking for RAG', level=3)
add_normal(
    'Each successfully parsed document is fed into chunk_document(). This function breaks '
    'the text into chunks suitable for vector embedding and retrieval-augmented generation.'
)
add_normal('The chunking strategy is section-aware:')
add_bullet('If the document has sections (headings, sheets, slides), each section is chunked independently.')
add_bullet('Within a section, text is split on paragraph boundaries, keeping each chunk under ~800 tokens (roughly 3,200 characters).')
add_bullet('If a section is small enough, it becomes a single chunk.')
add_bullet('If a document has no sections, a sliding-window fallback is used with a 100-token overlap.')

add_normal(
    'Each chunk gets a SHA-256-based ID (first 12 hex chars), a source file reference, '
    'the section title, a chunk index, and a token estimate. Chunks are saved to output/chunks.jsonl, '
    'one JSON object per line.'
)

# Step 4
doc.add_heading('Step 4 — Audit Checklist Generation (Local)', level=3)
add_normal(
    'Before calling any AI, main.py runs a local keyword-based audit. The generate_audit_checklist() '
    'function checks each parsed document\'s text against a predefined set of audit requirements. '
    'Each requirement has a list of keywords — for example, the BRD requirement looks for '
    '"brd", "business requirement", "requirement" in the document text.'
)
add_normal(
    'Matches produce an evidence snippet (a short excerpt around the keyword match). '
    'The results are written to an Excel workbook (Client_Audit_Checklist.xlsx) with three sheets: '
    'Document Evidence Matrix, Parsed Files Summary, and Gap Analysis.'
)

# Step 5
doc.add_heading('Step 5 — AI Analysis (Gemini)', level=3)
add_normal(
    'The analyzer module (analyzer.py) constructs a large prompt containing all parsed document '
    'text (capped at 10,000 characters per document) and the full SDS stage-gate standard as JSON. '
    'This prompt is sent to Google Gemini (model: gemini-2.5-flash-lite) with temperature=0.1 '
    'and response_mime_type="application/json" to force structured output.'
)
add_normal(
    'Gemini returns a single JSON object containing: client overview, document summaries, '
    'a workflow narrative, a 22-item checklist (mapped to the SDS standard), gap analysis, '
    'action items, readiness assessment, and metrics. The analyzer then post-processes this — '
    'coercing metric types, recalculating zeros, and computing Go/No-Go gate decisions using a '
    'deterministic decision engine (not another AI call).'
)

# Step 6
doc.add_heading('Step 6 — Cross-Document Intelligence', level=3)
add_normal(
    'cross_document_intel.py runs a two-phase analysis. Phase 1 uses regex to extract entities '
    '(field names, dates, metrics) from each document locally. Phase 1.5 compares these entities '
    'across documents to find conflicts (e.g., same field called "member_id" in one doc and '
    '"memberID" in another). Phase 2 sends all documents plus the local findings to Gemini for '
    'deeper cross-reference analysis. The results are merged and attached to the main JSON as '
    'analysis_output.cross_document_intel.'
)

# Step 7
doc.add_heading('Step 7 — Project Intelligence Engine', level=3)
add_normal(
    'project_intelligence.py is the final synthesis step. It computes multi-dimensional scores '
    '(Completeness, Quality, Consistency, Freshness) for each checklist item, aggregates them '
    'into Stage Readiness Scores per gate, generates Go/No-Go gate decisions from these scores, '
    'runs a proactive recommendation engine, and calls Gemini one final time to produce the '
    '"Project Mind" — a narrative synthesis of the entire project.'
)

# Step 8
doc.add_heading('Step 8 — Output & Frontend Delivery', level=3)
add_normal(
    'After each major step, the analysis JSON is written to two locations:'
)
add_bullet('red_duke_demo/output/analysis_output.json — backend archive')
add_bullet('frontend/public/analysis_output.json — where the React app fetches it')

add_normal(
    'An Excel report (Red_Duke_Report.xlsx) is also generated and copied to frontend/public/ '
    'for download. If SharePoint is configured, the output directory contents are uploaded back '
    'to a SharePoint folder.'
)

add_normal(
    'The React frontend (built with Vite) loads analysis_output.json via a fetch() call in App.jsx. '
    'The JSON is stored in a React Context (AnalysisContext) and consumed by all page components — '
    'Dashboard, Documents, Checklist, Gap Analysis, Gate Decisions, Cross-Doc Intel, Project Mind, '
    'and Workflow Narrative. Each page reads the fields it needs from context and renders them.'
)

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════
# PART 1: DATA FLOW — NON-DEVELOPER VERSION
# ═══════════════════════════════════════════════════════════════════════

doc.add_heading('Part 1: Data Flow', level=1)
doc.add_heading('Non-Developer Version', level=2)

add_blue(
    'Red Duke works like a document analyst that reads through an entire client folder '
    'and produces a structured report. Here is how data moves through the system from '
    'start to finish.'
)

doc.add_heading('1. Pulling Documents from SharePoint', level=3)
add_blue(
    'The system connects to a client\'s SharePoint site using secure credentials. '
    'It downloads every file in the designated project folder — Word documents, Excel workbooks, '
    'PowerPoint decks, Visio diagrams, PDFs, and more. If SharePoint is not available '
    '(for example, during a demo), it uses a local copy of the files instead.'
)

doc.add_heading('2. Reading and Understanding Each File', level=3)
add_blue(
    'Each downloaded file is opened and read by a specialized reader. Word documents '
    'are broken into sections based on their headings. Excel files are read sheet by sheet. '
    'PowerPoint slides are read including speaker notes. Visio workflow diagrams are parsed '
    'to extract the flow of process steps. The system captures the text, any tables, and '
    'metadata like author name and creation date.'
)

doc.add_heading('3. Breaking Documents into Smaller Pieces', level=3)
add_blue(
    'Large documents are split into smaller, manageable pieces (we call them "chunks"). '
    'This is done intelligently — the system respects section boundaries so that a chunk '
    'about "Data Mapping" stays together rather than getting split mid-sentence. These chunks '
    'are stored so the system can quickly find relevant information later.'
)

doc.add_heading('4. Initial Audit (Keyword Matching)', level=3)
add_blue(
    'Before any AI is involved, the system does a quick scan of every document looking for '
    'keywords that indicate required deliverables. For example, it looks for words like '
    '"business requirement", "project plan", "workflow", "UAT", "go-live" to check which '
    'SDS deliverables are present. This produces a first-pass audit checklist saved as an '
    'Excel file.'
)

doc.add_heading('5. AI Deep Analysis', level=3)
add_blue(
    'The text from all documents is sent to Google\'s Gemini AI along with the SDS '
    'stage-gate standard (the list of 22 deliverables required across 5 project stages). '
    'Gemini reads through everything and produces a detailed report: what each document '
    'contains, which deliverables have evidence, which are missing, what the gaps are, '
    'and what actions to take. It also writes a plain-English narrative of the project\'s '
    'workflow.'
)

doc.add_heading('6. Cross-Document Checks', level=3)
add_blue(
    'The system then checks for inconsistencies between documents. For instance, if one '
    'document says the go-live date is March 15 and another says April 1, that gets flagged. '
    'If a field is called "member_id" in one place and "MemberID" in another, that is caught too. '
    'This helps ensure all documents are aligned before they reach production.'
)

doc.add_heading('7. Scoring and Decisions', level=3)
add_blue(
    'Each deliverable gets scored on four dimensions: how complete it is, the quality of '
    'evidence, how consistent it is across documents, and how current it looks. These scores '
    'roll up into a readiness score for each of the five SDS stage gates (Initiation, Design, '
    'Implementation, Test, Deploy). The system then makes a Go / Conditional / No-Go '
    'recommendation for each gate.'
)

doc.add_heading('8. Displaying Results in the Dashboard', level=3)
add_blue(
    'All of this data is saved as a JSON file and placed where the web application can '
    'read it. When you open the Red Duke dashboard in your browser, it loads this file '
    'and displays everything across several pages: an overview dashboard, document details, '
    'the checklist, gap analysis, gate decisions, cross-document issues, and the "Project Mind" '
    '— a narrative summary of where the project stands.'
)

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════
# PART 2: ANALYSIS ENGINE — DEVELOPER VERSION
# ═══════════════════════════════════════════════════════════════════════

doc.add_heading('Part 2: Analysis Engine', level=1)
doc.add_heading('Developer Version', level=2)

add_normal(
    'This section explains exactly how Red Duke analyzes documents — the prompts, keywords, '
    'scoring logic, and guardrails that power each analysis step.'
)

# SDS Standard
doc.add_heading('The SDS Stage-Gate Standard', level=3)
add_normal(
    'Everything in Red Duke maps back to the SDS implementation lifecycle. We define 22 '
    'deliverables across 5 stage gates in the SDS_STANDARD dict (analyzer.py). Each deliverable '
    'has a name and a responsible role:'
)
add_bullet('TG-1 (Initiation) — 5 deliverables: BRD, Project Plan, Kickoff Presentation, Security Docs, Training Docs')
add_bullet('TG0 (Design) — 4 deliverables: Workflow Diagram, Data Mapping, Tech Design Spec, BRD Sign-off')
add_bullet('TG1 (Implementation) — 5 deliverables: Transition Workbook, Config Docs, SLA Docs, Mailroom Config, Eligibility/Claims Config')
add_bullet('TG2 (Test) — 4 deliverables: UAT Plan, UAT Sign-off, TPM/Compliance Checklist, Functional Testing Evidence')
add_bullet('TG3 (Deploy/Warranty) — 4 deliverables: M2P/Go-Live Sign-off, Warranty Checklist, Lessons Learned, Final Handoff')

add_normal(
    'This standard is injected as JSON directly into the Gemini prompt so the AI knows exactly '
    'what to look for and how to categorize its findings.'
)

# Keyword-based audit
doc.add_heading('Local Keyword Matching (Step 4)', level=3)
add_normal(
    'Before the AI runs, we do a deterministic keyword scan. Each audit requirement has a keywords list. '
    'The system converts each document\'s full text to lowercase and checks for substring matches. '
    'Here are the actual keyword sets we use:'
)

keywords_table = doc.add_table(rows=13, cols=2, style='Table Grid')
keywords_table.cell(0, 0).text = 'Deliverable'
keywords_table.cell(0, 1).text = 'Keywords'
kw_data = [
    ('BRD', 'brd, business requirement, requirement'),
    ('Project Plan', 'project plan, critical path, roadmap, timeline, initiation'),
    ('Kickoff Presentation', 'kickoff, scope review, agenda, introductions'),
    ('Workflow Diagram', 'workflow, diagram, process flow, data capture, image import'),
    ('Transition Workbook', 'transition, workbook, implementation, configuration, sop'),
    ('SLA Documentation', 'report, reporting, sla, volume, frequency'),
    ('Testing / UAT', 'testing, uat, test case, functional testing, volume testing'),
    ('Security Docs', 'security, access control, sso, single sign-on, password'),
    ('Eligibility/Claims', 'eligibility, claim, matching, reject, routing, 837'),
    ('M2P / Go-Live', 'm2p, go-live, move to production, sign-off, warranty'),
    ('TPM / Compliance', 'tpm, checklist, compliance, complete?'),
    ('Training Docs', 'training, process awareness, onboarding'),
]
for i, (deliv, kws) in enumerate(kw_data):
    keywords_table.cell(i + 1, 0).text = deliv
    keywords_table.cell(i + 1, 1).text = kws

# Format table header
for cell in keywords_table.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.size = Pt(10)

add_normal(
    'When a keyword is found, the system extracts a short snippet (30 characters before '
    'and 60 after the match) as evidence. Each deliverable is marked FOUND or MISSING.'
)

# The Gemini prompt
doc.add_heading('The Gemini Analysis Prompt (Step 5)', level=3)
add_normal(
    'The core analysis prompt in analyzer.py is structured as follows. This is the actual '
    'prompt structure sent to the AI:'
)
add_normal('1. System role definition:')
add_code_block(
    '"You are RED DUKE, an expert implementation documentation analyst at\n'
    'Smart Data Solutions (SDS). SDS is a healthcare technology company that\n'
    'automates complex payer operations — digital mailroom, document intelligence,\n'
    'claims processing, and interoperability — for 500+ healthcare organizations."'
)

add_normal('2. Mission instructions — the AI is told to:')
add_bullet('Deeply understand what each document contains, noting placeholders and TBDs')
add_bullet('Cross-reference documents against each other')
add_bullet('Map evidence to the SDS stage-gate standard')
add_bullet('Identify gaps and prioritize by business risk')
add_bullet('Produce a workflow narrative a new PM could use on day one')

add_normal('3. The full SDS_STANDARD is injected as formatted JSON.')

add_normal(
    '4. All document text is appended, each wrapped with a header block showing the '
    'filename, type, size, section count, and table count. Document text is capped at '
    '10,000 characters each to control prompt size and cost.'
)

add_normal('5. The expected JSON response schema is spelled out field by field. Key fields include:')
add_bullet('client_overview — name, summary, scope, integrations, risks, current phase')
add_bullet('documents[] — per-file purpose, summary, key topics, artifacts present/missing, quality score (0-100)')
add_bullet('checklist[] — all 22 deliverables with status (PRESENT/PARTIAL/MISSING), evidence files, confidence level, evidence trail with verbatim excerpts, priority (P1/P2/P3), estimated effort, is_blocker flag')
add_bullet('gap_analysis[] — each gap with impact (HIGH/MEDIUM/LOW), recommendation, responsible party, effort, dependencies')
add_bullet('action_items[] — prioritized actions with owner, gate, deliverable, effort')
add_bullet('readiness_assessment — overall RED/YELLOW/GREEN, next gate, blockers, quick wins, executive summary')
add_bullet('metrics — coverage score, gap counts, per-gate completion percentages')

# Guardrails
doc.add_heading('Guardrails and Quality Controls', level=3)
add_normal('We enforce several guardrails on the AI output:')

add_bullet('Temperature is set to 0.1 — this makes Gemini\'s output nearly deterministic and reduces hallucination.')
add_bullet('response_mime_type is set to "application/json" — Gemini is forced to return valid JSON, not prose.')
add_bullet('Markdown fence stripping — if Gemini wraps its response in ```json fences (a common habit), we strip them before parsing.')
add_bullet('JSONDecoder.raw_decode() — if Gemini returns multiple JSON objects concatenated, we only parse the first valid one instead of crashing.')
add_bullet('Type coercion — all metric values are cast to int after parsing, since Gemini sometimes returns numbers as strings.')
add_bullet('Zero recalculation — if Gemini returns 0 for coverage_score or gap counts, we recalculate them ourselves from the checklist data.')
add_bullet('Stage-gate completion recomputation — if all gate percentages come back as 0, we recompute them from the checklist by matching deliverable names to the SDS_STANDARD.')
add_bullet('The checklist MUST include all 22 deliverables — the prompt explicitly states this count.')
add_bullet('Evidence trail requirement — for every PRESENT or PARTIAL item, the AI must provide at least one evidence trail entry with a verbatim excerpt from the source document.')

# Go/No-Go Engine
doc.add_heading('Go/No-Go Decision Engine (Deterministic)', level=3)
add_normal(
    'Gate decisions are NOT made by AI. After Gemini returns the checklist and gap analysis, '
    'a deterministic Python function (_compute_gate_decisions in analyzer.py) computes the '
    'decision for each gate. The logic is:'
)
add_bullet('GO — all blocker deliverables are PRESENT, confidence >= 70%, no HIGH-impact gaps for this gate')
add_bullet('CONDITIONAL — some blockers are PARTIAL, or confidence is between 50-69%, or non-critical gaps remain')
add_bullet('NO-GO — any blocker is MISSING, confidence < 50%, or HIGH-impact gaps exist')

add_normal(
    'Each decision includes a confidence score (weighted average of item confidence levels), '
    'completion percentage, blocker list with status, optional gaps, high-risk gaps, '
    'estimated delay, and a human-readable rationale string.'
)

add_normal(
    'In the Project Intelligence module, a second gate decision engine '
    '(generate_gate_decisions in project_intelligence.py) runs on top of the '
    'multi-dimensional scores. This one uses readiness thresholds: >= 70% for GO, '
    '>= 40% for CONDITIONAL, below 40% for NO-GO.'
)

# Cross-document analysis
doc.add_heading('Cross-Document Intelligence (Step 6)', level=3)
add_normal('The cross-document engine in cross_document_intel.py works in two phases:')

add_normal('Phase 1 — Local entity extraction using regex:')
add_bullet('Field names — matched via patterns for snake_case, camelCase, PascalCase, and UPPER_CASE identifiers (4-40 characters)')
add_bullet('Dates — ISO format (2025-03-15), US format (03/15/2025), and written format (March 15, 2025)')
add_bullet('Metrics/SLAs — numbers followed by units like %, days, hours, minutes, SLA, business days')

add_normal(
    'Extracted fields are normalized (lowercased, separators stripped) and looked up against '
    'a synonym table. For example, "member_id", "memberID", "MemberId", "mbr_id", and "MbrId" '
    'all map to the canonical form "member_id". The synonym groups cover healthcare-specific '
    'terms: member_id, provider_id, claim_id, date_of_service, date_of_birth, group_id, '
    'subscriber_id, tax_id, go_live, and uat.'
)

add_normal('Phase 1.5 — Local conflict detection:')
add_bullet('Field mismatches — if the same canonical field appears with different raw names in different documents, that is flagged as HIGH severity.')
add_bullet('Metric conflicts — if the same type of metric (e.g., SLA turnaround) has different numeric values across documents, that is flagged.')

add_normal('Phase 2 — Gemini deep analysis:')
add_normal(
    'A second prompt is sent to Gemini containing all document text (capped at 8,000 chars '
    'per doc), the extracted entities, and the local conflicts found so far. The prompt '
    'instructs Gemini to find NEW issues (not repeat the local ones) across five categories:'
)
add_bullet('Field/Schema Mismatches')
add_bullet('Timeline Conflicts')
add_bullet('Terminology Drift')
add_bullet('Requirements vs. Implementation Gaps')
add_bullet('Numeric/Metric Conflicts')

add_normal(
    'Gemini returns a JSON with cross_document_issues, a cross_reference_matrix (how each '
    'pair of documents relates), and a consistency score. Local and AI issues are merged, '
    'severity counts are recalculated, and the consistency score is computed as: '
    '100 - (HIGH_issues * 15) - (MEDIUM_issues * 8) - (LOW_issues * 3).'
)

# Multi-dimensional scoring
doc.add_heading('Multi-Dimensional Scoring (Step 7)', level=3)
add_normal(
    'The Project Intelligence module scores every checklist item on four dimensions:'
)

scoring_table = doc.add_table(rows=5, cols=3, style='Table Grid')
scoring_table.cell(0, 0).text = 'Dimension'
scoring_table.cell(0, 1).text = 'Weight'
scoring_table.cell(0, 2).text = 'How It Is Calculated'
scoring_data = [
    ('Completeness', '40%', 'Based on status: PRESENT starts at 75 (with bonuses for rich evidence trails and multiple evidence files up to 100), PARTIAL starts at 35 (capped at 65), MISSING = 0.'),
    ('Quality', '25%', 'Based on confidence level (HIGH=80, MEDIUM=55, LOW=25), blended 60/40 with source document quality scores (STRONG=90, ADEQUATE=60, WEAK=25). Bonus for evidence trails with excerpts and assessments.'),
    ('Consistency', '20%', 'Starts at 95 if no cross-document issues affect the evidence files. Deducts 25 per HIGH issue, 12 per MEDIUM, 5 per LOW that references the same files. 0 if MISSING, 50 if no evidence files to check.'),
    ('Freshness', '15%', 'Penalizes draft content — scans evidence trails for signals like "TBD", "placeholder", "draft", "pending", "WIP". If found, score drops (60 minus 15 per signal). Otherwise based on document completeness (section/table counts).'),
]
for i, (dim, weight, calc) in enumerate(scoring_data):
    scoring_table.cell(i + 1, 0).text = dim
    scoring_table.cell(i + 1, 1).text = weight
    scoring_table.cell(i + 1, 2).text = calc

for cell in scoring_table.rows[0].cells:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.size = Pt(10)

add_normal(
    'The Readiness Score for each item is the weighted composite: '
    'Completeness×0.40 + Quality×0.25 + Consistency×0.20 + Freshness×0.15. '
    'Stage Readiness is the average of all item readiness scores within that gate. '
    'Overall Readiness is the average across all gates that have at least one item.'
)

# Proactive recommendations
doc.add_heading('Proactive Recommendation Engine', level=3)
add_normal('The recommendation engine ranks actions by impact score:')
add_bullet('MISSING deliverables start at 100 if they are blockers, 60 otherwise. P1 priority adds 30, P2 adds 15. Small effort adds 15, medium adds 10, large adds 5.')
add_bullet('PARTIAL deliverables start at 80 (blocker) or 50 (non-blocker), plus 30% of the gap to 100% readiness.')
add_bullet('Cross-document issues are scored by severity: HIGH=85, MEDIUM=55, LOW=30.')
add_normal('All candidates are sorted by impact score descending and the top 10 are returned with rank numbers.')

# Project Mind
doc.add_heading('Project Mind (Final AI Synthesis)', level=3)
add_normal(
    'The last Gemini call produces the "Project Mind" — a narrative synthesis. The prompt '
    'provides Gemini with all computed data: client overview, deliverable counts, gap counts, '
    'cross-document issue counts, consistency score, and gate decisions. Gemini returns:'
)
add_bullet('project_understanding — 3-4 paragraphs, the system\'s coherent understanding of the project')
add_bullet('critical_insight — one sentence for leadership')
add_bullet('risk_narrative — the biggest risk and why it matters')
add_bullet('momentum_assessment — ACCELERATING / STEADY / STALLING / BLOCKED')
add_bullet('predicted_bottleneck — the specific deliverable most likely to cause delay')
add_bullet('time_to_next_gate — estimated calendar time')
add_bullet('next_actions_narrative — what the PM should do tomorrow morning')

add_normal('Temperature for this call is 0.15 — slightly higher than the analysis call to allow more natural language, but still very controlled.')

doc.add_page_break()


# ═══════════════════════════════════════════════════════════════════════
# PART 2: ANALYSIS ENGINE — NON-DEVELOPER VERSION
# ═══════════════════════════════════════════════════════════════════════

doc.add_heading('Part 2: Analysis Engine', level=1)
doc.add_heading('Non-Developer Version', level=2)

add_blue(
    'This section explains how Red Duke actually analyzes documents and produces '
    'its recommendations, written for a non-technical audience.'
)

doc.add_heading('The SDS Standard It Checks Against', level=3)
add_blue(
    'SDS has a standard implementation lifecycle with five stages (called "stage gates"): '
    'Initiation, Design, Implementation, Test, and Deploy/Warranty. Each stage requires '
    'certain documents to be completed — 22 deliverables in total. For example, Initiation '
    'requires a Business Requirements Document and a Project Plan, while the Test stage '
    'requires UAT test cases and sign-off. Red Duke knows this standard and checks every '
    'client folder against it.'
)

doc.add_heading('How It Reads Documents', level=3)
add_blue(
    'Each file type is read differently. Word documents are read paragraph by paragraph, '
    'with headings used to identify sections. Excel files are read sheet by sheet, row by row. '
    'PowerPoint files are read slide by slide, including speaker notes. Visio diagrams '
    '(workflow charts) are opened and each shape\'s text is extracted to understand the '
    'process flow. The system captures everything — headings, body text, tables, notes, '
    'metadata.'
)

doc.add_heading('The First Check: Keyword Scanning', level=3)
add_blue(
    'Before involving AI, the system does a straightforward scan. It looks for specific '
    'words in each document that indicate whether a required deliverable exists. For example, '
    'if it finds "business requirement" or "BRD" anywhere in the documents, it marks the '
    'Business Requirements Document as found. If it finds "UAT" or "test case", it marks '
    'testing evidence as found. This is a fast, reliable first pass.'
)

doc.add_heading('The AI Analysis', level=3)
add_blue(
    'The text from all documents is then sent to Google\'s Gemini AI. The AI is given '
    'the full SDS standard (all 22 required deliverables) and asked to do a thorough '
    'review. It reads through everything and determines:'
)
p = doc.add_paragraph()
run = p.add_run(
    '— What each document is about and how complete it is\n'
    '— Which of the 22 required deliverables have evidence, which are partially done, '
    'and which are missing entirely\n'
    '— What the specific gaps are and how risky each gap is\n'
    '— What actions need to happen next and who should own them\n'
    '— A plain-English narrative of the project\'s workflow'
)
run.font.color.rgb = RGBColor(0x00, 0x50, 0xA0)

add_blue(
    'The AI is set to be very precise (low "temperature" setting) and is required to return '
    'structured data rather than free-form text. For every finding, it must cite the specific '
    'document and quote the exact text that supports its conclusion. This makes every result '
    'traceable back to a real document.'
)

doc.add_heading('Cross-Document Consistency Checks', level=3)
add_blue(
    'After the main analysis, the system checks whether documents agree with each other. '
    'It looks for five types of problems:'
)
p = doc.add_paragraph()
run = p.add_run(
    '— Field name mismatches (e.g., "member_id" vs "MemberID" — same thing, different names)\n'
    '— Timeline conflicts (different dates for the same milestone)\n'
    '— Terminology drift (same concept called different names across documents)\n'
    '— Requirements vs. reality gaps (something required in one doc but missing in another)\n'
    '— Number conflicts (SLA or volume numbers that do not match across documents)'
)
run.font.color.rgb = RGBColor(0x00, 0x50, 0xA0)

add_blue(
    'The system knows common healthcare field synonyms — it understands that "DOS" and '
    '"date_of_service" refer to the same thing, for instance. Each inconsistency gets a '
    'severity rating (High, Medium, or Low) and a recommendation for how to fix it.'
)

doc.add_heading('Scoring Each Deliverable', level=3)
add_blue(
    'Each of the 22 deliverables is scored on four dimensions. Completeness measures how '
    'much of the deliverable exists. Quality measures how thorough and reliable the evidence '
    'is. Consistency measures whether the evidence agrees across documents. Freshness checks '
    'whether the content looks final or still has placeholders like "TBD" and "draft."'
)
add_blue(
    'These four scores are combined (with Completeness weighted most heavily) into a single '
    'Readiness Score for each deliverable. The scores for all deliverables within a stage '
    'gate are averaged to produce a Stage Readiness Score. All stage scores together give '
    'the Overall Project Readiness percentage.'
)

doc.add_heading('Go / No-Go Decisions', level=3)
add_blue(
    'For each of the five stage gates, the system makes a recommendation:'
)
p = doc.add_paragraph()
run = p.add_run(
    '— GO means all required items are in place and the project can advance through this gate\n'
    '— CONDITIONAL means most things are ready but a few items still need to be completed\n'
    '— NO-GO means critical items are missing and the project cannot advance until they are addressed'
)
run.font.color.rgb = RGBColor(0x00, 0x50, 0xA0)

add_blue(
    'These decisions are made by rules, not by AI — they are based directly on the scores '
    'and whether blocking deliverables are present. This makes the decisions predictable '
    'and explainable.'
)

doc.add_heading('Prioritized Recommendations', level=3)
add_blue(
    'Finally, the system ranks the most impactful actions that should be taken. Missing '
    'blockers get the highest priority. Items that are partially complete get scored based '
    'on how close they are to being done. Cross-document inconsistencies are also ranked. '
    'The top 10 actions are presented in order of impact, with effort estimates and the '
    'person responsible.'
)

doc.add_heading('The "Project Mind"', level=3)
add_blue(
    'The last step produces a narrative summary — the system\'s understanding of where '
    'the project stands, what the biggest risk is, whether the project has momentum or is '
    'stalling, what the predicted bottleneck will be, and what the PM should do first thing '
    'tomorrow morning. This gives leadership a single-page view of project health without '
    'having to read through every detail.'
)

doc.add_heading('Safety and Accuracy Measures', level=3)
add_blue(
    'Several safeguards are in place to ensure accuracy:'
)
p = doc.add_paragraph()
run = p.add_run(
    '— The AI is set to be highly precise, minimizing creative interpretation\n'
    '— Every finding must include a direct quote from the source document\n'
    '— If the AI returns incomplete numbers, the system recalculates them independently\n'
    '— Gate decisions are rule-based, not AI-generated, so they are consistent and predictable\n'
    '— The system checks for and handles common AI output quirks (formatting issues, '
    'type mismatches, missing data)\n'
    '— If any single file fails to parse, the rest of the pipeline continues unaffected'
)
run.font.color.rgb = RGBColor(0x00, 0x50, 0xA0)


# ── Save ──
output_path = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    'red_duke_demo', 'output', 'Red_Duke_Technical_Documentation.docx'
)
os.makedirs(os.path.dirname(output_path), exist_ok=True)
doc.save(output_path)
print(f"Document saved to: {output_path}")
